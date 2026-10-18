"""Builds presentation/TourPlan_AI_Presentation.pptx (15 slides, speaker notes)."""
import pathlib
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import content as c

ROOT = c.ROOT
D = ROOT / "diagrams"
BG, CARD, TXT, AMB = RGBColor(0x0B, 0x1F, 0x1A), RGBColor(0x14, 0x33, 0x2B), RGBColor(0xF3, 0xEF, 0xE2), RGBColor(0xF5, 0xB0, 0x41)
E = c.EXAMPLE
DEMO = c.DEMO
tests = c.tests()
eps = c.endpoints()

# ---- illustrative chart from the real budget module (excluding attraction fees)
cats = {k: v for k, v in DEMO.items() if isinstance(v, (int, float)) and k not in ("total", "per_person", "per_day", "core_total", "rooms")}
cats = {k: v for k, v in DEMO.get("categories", cats).items() if v}
fig, ax = plt.subplots(figsize=(5.2, 4.2), facecolor="#0B1F1A")
ax.pie(list(cats.values()), labels=[k.replace("_", " ") for k in cats], wedgeprops=dict(width=.42),
       textprops=dict(color="#F3EFE2", fontsize=9), colors=["#2f855a", "#f5b041", "#e76f51", "#4fd1c5", "#a0aec0", "#d69e2e", "#68d391"])
ax.set_title(f"Illustrative estimate: Mumbai to Lonavala, 4 days, 3 people\n(sample data, excl. attraction fees) total Rs {DEMO['total']:,.0f}",
             color="#F3EFE2", fontsize=9)
chart = pathlib.Path(__file__).parent / "demo_budget.png"
fig.savefig(chart, dpi=160, facecolor=fig.get_facecolor(), bbox_inches="tight")

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
blank = prs.slide_layouts[6]


def slide(title, bullets=None, img=None, notes="", img_w=6.2, sub=None):
    s = prs.slides.add_slide(blank)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    tb = s.shapes.add_textbox(Inches(.6), Inches(.35), Inches(12), Inches(.9)).text_frame
    tb.text = title
    r = tb.paragraphs[0].runs[0]
    r.font.size, r.font.bold, r.font.color.rgb = Pt(32), True, AMB
    tb.paragraphs[0].alignment = PP_ALIGN.LEFT
    if sub:
        p = tb.add_paragraph(); p.text = sub; p.runs[0].font.size = Pt(16); p.runs[0].font.color.rgb = TXT
    if bullets:
        w = 12 if not img else 12.1 - img_w - .3
        box = s.shapes.add_shape(1, Inches(.6), Inches(1.5), Inches(w), Inches(5.3))
        box.fill.solid(); box.fill.fore_color.rgb = CARD; box.line.fill.background()
        tf = box.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = tf.margin_top = Inches(.25)
        for i, b in enumerate(bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = "• " + b
            p.runs[0].font.size, p.runs[0].font.color.rgb = Pt(19), TXT
            p.space_after = Pt(10); p.alignment = PP_ALIGN.LEFT
    if img:
        from PIL import Image
        iw, ih = Image.open(img).size
        W = img_w if bullets else 11
        H = W * ih / iw
        if H > 5.4:
            H = 5.4; W = H * iw / ih
        s.shapes.add_picture(str(img), Inches(13.333 - .6 - W) if bullets else Inches((13.333 - W) / 2), Inches(1.5), Inches(W), Inches(H))
    s.notes_slide.notes_text_frame.text = notes
    return s


slide(c.PROJECT, ["Software Engineering Lab project, MBA Tech Computer Engineering",
                  f"Presented by: {c.STUDENT}  ({c.ROLL})", f"Faculty mentor: {c.FACULTY}", f"{c.INSTITUTE}   |   {c.DATE}",
                  "Placeholders in square brackets must be filled in by the student."],
      notes="Introduce yourself, the project name and mentor. Say in one sentence: it plans trips and is honest about estimates.")
slide("Problem and motivation", ["Choosing where to go within a budget is hard", "Costs are scattered across sites and spreadsheets",
                                 "Generic chatbots can invent prices and places", "Need: transparent, reproducible, editable plans"],
      notes="Viva: Why not just use a chatbot? Because numbers must be reproducible and testable; LLMs may hallucinate.")
slide("Objectives and scope", ["Transparent 100-point destination ranking", "Itinerary, budget, map, packing list, PDF", "AI assistant with honest fallback",
                               "Security: hashing, JWT, per-user access, no keys in browser", "Out of scope: booking, live prices, payments"],
      notes="State clearly the app does not book anything and prices are indicative estimates.")
slide("Existing vs proposed system", ["Existing: booking portals (live prices, no planning), spreadsheets (manual), chatbots (may hallucinate)",
                                      "Proposed: one app combining discovery, planning, budgeting, map and assistant",
                                      "LLM optional and limited to explanation; labelled fallback"],
      notes="Emphasise the design choice: engine decides, LLM explains.")
slide("System architecture", ["React SPA", "FastAPI REST + JWT", "SQLAlchemy + SQLite", "Deterministic engine", "Optional Gemini, Open-Meteo, Frankfurter"],
      img=D / "architecture.png", notes="Walk the diagram left to right. The key is only in the backend environment.", img_w=7)
slide("Use cases", None, img=D / "use_case.png", notes="Actors: guest and registered user. Highlight Generate trip plan and Chat with assistant.")
slide("Data flow and activity", None, img=D / "activity_trip_planning.png",
      notes="Show validation, scoring, budget check and the over-budget branch that suggests cheaper alternatives.")
slide("Database design", [f"{c.N_DEST} destinations, {c.N_ATTR} attractions (sample data)", "Trips own days, activities, expenses, packing items",
                          "Unique constraints: trip/day and user/destination", "Foreign keys enabled in SQLite"],
      img=D / "er_diagram.png", img_w=7, notes="Explain one-to-many relationships. Seed data is hand-written sample data.")
slide("AI architecture", ["1 Rule-based requirement extraction", "2 Deterministic recommender, planner, budget", "3 Optional Gemini explanation as validated JSON",
                          "4 Fallback labelled rule_based", "LLM never supplies prices or coordinates"],
      img=D / "sequence_ai_generation.png", img_w=6.4, notes="Viva: what if there is no key? Rule-based text, labelled honestly. Rate limit 10 per minute by default.")
slide("Recommendation score", [f"{k.replace('_', ' ')}: {v} points" for k, v in c.WEIGHTS.items()] +
      [f"+ up to {c.HINT_BONUS} bonus for place-type hints", "In-budget destinations rank first; every component shows a reason"],
      notes="Total is 100 points. Show the breakdown in the UI during the demo.")
slide("Budget estimation", [f"Worked example total Rs {E['total']:,.0f} (matches unit test)", "Categories incl. 8% misc and 10% emergency buffer",
                            "Rooms = ceil(travellers / 2)", "All figures are estimates from sample data"],
      img=chart, img_w=5.2, notes="The chart is computed by the real budget module for Mumbai to Lonavala, 3 people, 4 days, train, excluding attraction fees. It is illustrative.")
slide("User interface", ["Dark forest glassmorphism; featured card, stat rings, map panel", "Responsive; keyboard focus; aria roles; reduced motion",
                        "Airplane transition with a settings toggle", "[Insert real screenshots here after running the app]"],
      notes="Replace the last bullet with screenshots captured from the running application. None were faked.")
slide("Security", ["bcrypt password hashing", "JWT with expiry; protected routes", "Other users' trips return 404", "Input validation, CORS allow-list, AI rate limit",
                   "Limitation: token in localStorage"], notes="Viva: XSS risk of localStorage; cookies with CSRF protection are future work.")
slide("Testing", [f"{len(tests)} pytest functions in the repository", "Budget and travel as pure functions with hand-calculated values",
                  "API tests with temp DB, LLM disabled", "Results NOT yet recorded: run pytest and npm test, then insert real output here"],
      notes="Be honest: do not claim a pass count until you have run the tests. Paste real output in report section 25.")
slide("Demo, limitations and future work", ["Demo: Mumbai, 3 people, 4 days, Rs 20,000; ranking, itinerary edit, budget, map, chat, PDF",
                                            "Limits: sample data, 10 start cities, rule-based extraction, in-memory rate limit",
                                            "Future: migrations, live pricing, cookie sessions, larger dataset", "Thank you. Questions?"],
      notes="DEMO STEPS: 1 Discover and filter. 2 Log in demo@tourplan.ai / demo12345. 3 New trip. 4 Score breakdown. 5 Edit and save itinerary. "
            "6 Budget warning and alternatives. 7 Map route. 8 Assistant: 'Plan 3 days in Lonavala for 3 people under 15k' then 'make it cheaper'. "
            "9 PDF export. 10 Turn transition off in Settings. VIVA: see docs/VIVA_AND_DEMO.md for 20 prepared questions.")
(ROOT / "presentation").mkdir(exist_ok=True)
prs.save(ROOT / "presentation/TourPlan_AI_Presentation.pptx")
print("slides:", len(prs.slides._sldIdLst))
