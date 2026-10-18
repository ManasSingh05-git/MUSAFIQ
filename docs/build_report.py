"""Builds report/TourPlan_AI_Report.docx and .pdf from one section structure (31 sections)."""
import pathlib
import content as c

ROOT = c.ROOT
D = ROOT / "diagrams"
E = c.EXAMPLE
eps = c.endpoints()
tests = c.tests()
SHOTS = ["Landing page", "Dashboard", "Discover", "Destination detail", "Smart Trip Generator",
         "Itinerary", "Budget planner", "Interactive map", "My Trips", "AI assistant", "Profile / settings"]

S = []  # (title, [blocks])


def sec(title, *blocks):
    S.append((title, list(blocks)))


P = lambda t: ("p", t)
B = lambda *i: ("b", list(i))
T = lambda head, rows: ("t", head, rows)
I = lambda f, cap: ("img", D / f, cap)
PH = lambda t: ("ph", t)

sec("Cover Page", P(c.PROJECT), P("Software Engineering Lab Project Report"), P("MBA Tech, Computer Engineering"),
    PH(f"{c.STUDENT}   |   {c.ROLL}"), PH(f"Guided by: {c.FACULTY}"), PH(c.INSTITUTE), PH(c.DATE))
sec("Certificate", P("This is to certify that the project entitled \"" + c.PROJECT + "\" is a bonafide work carried out by the "
                     "student(s) named below in partial fulfilment of the Software Engineering Lab requirement."),
    PH(f"Student(s): {c.STUDENT}, {c.ROLL}"), PH("Faculty signature: ______________________ (to be signed by the faculty)"),
    PH("Head of Department signature: ______________________ (to be signed)"))
sec("Declaration", P("I/We declare that this work is our own, that sources are acknowledged in the references, and that "
                     "no results in this report have been invented. Sections marked as placeholders must be completed by the student."),
    PH(f"Name and signature: {c.STUDENT} ______________________  Date: {c.DATE}"))
sec("Acknowledgement", P("We thank our faculty mentor, the department and everyone who gave feedback on this project. "
                         "We also acknowledge the open-source communities behind FastAPI, React, Leaflet, OpenStreetMap and the other tools listed in the references."),
    PH("Edit this paragraph to add personal thanks."))
sec("Abstract",
    P("TourPlan AI is a web application that helps a traveller turn a rough idea (for example, 'three days near Mumbai under "
      "Rs 20,000') into a ranked destination list, a day-by-day itinerary, an itemised budget estimate, a map and a PDF. "
      "A deterministic engine produces every number and place from a sample dataset; an optional large language model "
      "(Gemini, configured only on the server) is limited to explaining the plan, and a rule-based fallback is used and "
      "labelled whenever the model is unavailable. The system uses React, FastAPI, SQLAlchemy and SQLite. "
      "This report documents the requirements, design, implementation and test approach. "
      "Test results are not stated here because the tests were not executed in the authoring environment; section 25 is a placeholder for real output."))
sec("Table of Contents", ("toc",))
sec("1. Introduction",
    P("Trip planning normally means juggling search engines, maps, spreadsheets and chat apps. TourPlan AI brings destination "
      "discovery, itinerary building, budgeting, maps and an assistant into one application, and is built as a Software "
      "Engineering Lab project to practise requirement analysis, design, implementation and testing."))
sec("2. Problem Statement",
    P("Travellers struggle to decide where to go within a budget, to estimate total cost, and to arrange a realistic daily schedule. "
      "Generic chatbots can invent prices and places. The problem is to build a planner that is transparent about how it recommends, "
      "keeps all figures reproducible, and is honest about the limits of its data."))
sec("3. Existing System",
    B("Booking portals show live prices but do not plan a whole trip or explain recommendations.",
      "Spreadsheets give control but need manual research and arithmetic.",
      "General chatbots are conversational but can hallucinate prices, places and distances, and give no editable plan."))
sec("4. Proposed System",
    P("A planner whose recommendation score and budget formulas are visible, whose places come from a database, and whose LLM is "
      "optional and constrained to explanation. Plans are saved, editable, mappable and exportable."))
sec("5. Objectives and Scope",
    B("Rank destinations transparently (100-point score with reasons).", "Generate and edit itineraries.",
      "Estimate and track budget with over-budget warnings.", "Provide map, packing list, PDF export and an assistant.",
      "Keep secrets on the server, validate inputs, protect user data."),
    P("Out of scope: booking, live prices, payments, real-time availability, multi-currency trips, mobile native apps."))
sec("6. Functional Requirements", T(["ID", "Requirement"], c.FR))
sec("7. Non-Functional Requirements", T(["ID", "Requirement"], c.NFR))
sec("8. Feasibility Study",
    B("Technical: all tools are free and open source; runs on a laptop.", "Economic: no paid services required; the LLM key is optional.",
      "Operational: single-user demo deployment; simple setup via two commands.", "Schedule: modular design allowed backend, frontend and documents to progress in parallel."))
sec("9. Technology Stack", T(["Layer", "Technology", "Why"], c.STACK))
sec("10. System Architecture",
    P("The browser runs a React single-page app. It calls the FastAPI REST API with a JWT. The API reads and writes SQLite via SQLAlchemy, "
      "calls the deterministic engine, and optionally calls Gemini and public weather/currency services."),
    I("architecture.png", "Figure 1: System architecture"))
sec("11. Use Case Design", I("use_case.png", "Figure 2: Use-case diagram"), T(["Use case", "Actor", "Description"], c.USE_CASES))
sec("12. Data Flow Diagrams", I("data_flow.png", "Figure 3: Data flow diagram"),
    I("activity_trip_planning.png", "Figure 4: Activity diagram, trip planning"))
sec("13. ER Diagram and Database Schema", I("er_diagram.png", "Figure 5: Entity-relationship diagram"),
    P("Tables: users, user_preferences, destinations, attractions, trips, itinerary_days, itinerary_activities, budget_expenses, "
      "favorites, chat_conversations, chat_messages, packing_items. Constraints include unique (trip, day_number), unique (user, destination) "
      f"favourites and enabled foreign keys. Seed data: {c.N_DEST} destinations and {c.N_ATTR} attractions (sample data)."))
sec("14. AI Architecture",
    P("Step 1: rule-based extraction turns chat text into structured requirements (city, days, travellers, budget, interests). "
      "Step 2: the recommender and planner build the plan from the database. Step 3: if a key is configured, Gemini is asked to return "
      "JSON explaining the plan; the reply is validated against a Pydantic model. Step 4: on any failure (no key, timeout, invalid JSON) "
      "a rule-based explanation is used and marked source = rule_based. The LLM never supplies prices, coordinates or place names."),
    I("sequence_ai_generation.png", "Figure 6: Sequence diagram, AI trip generation"),
    T(["Component", "Max points"], list(c.WEIGHTS.items())),
    P(f"Plus up to {c.HINT_BONUS} bonus points for place-type hints. Destinations within budget rank first; every component shows a reason."))
sec("15. User Interface Design",
    P("Dark forest-green glassmorphism with amber and coral accents; responsive grid; featured destination card, stat cards, map panel, "
      "progress rings; accessibility features include labelled controls, keyboard focus, aria roles and reduced-motion handling."),
    B("Airplane transition between routes, with a settings toggle and automatic reduced-motion bypass."))
sec("16. Screenshots", P("Screenshots must be captured from the running application. They are placeholders here; nothing has been faked."),
    *[PH(f"[Insert screenshot: {s}]") for s in SHOTS])
sec("17. Module Description",
    B("Authentication and profile", "Catalogue (destinations, favourites, stats, weather, currency)", "Trip generator (recommender, planner, transport, budget)",
      "Trips (CRUD, itinerary, budget, packing, PDF)", "AI assistant (extraction, rate limit, Gemini, fallback, conversations)",
      "Frontend pages and shared components"))
sec("18. API Design", P(f"{len(eps)} endpoints (extracted from the router source):"), T(["Method", "Path", "Handler"], eps))
sec("19. Budget Methodology",
    P("Categories: transport, accommodation, food, activities, local transport, shopping/misc (8% of core) and emergency buffer (10% of core). "
      "Rooms = ceil(travellers / 2). Style factors scale costs."),
    P(f"Worked example (identical to the unit test): 4 days, 3 travellers, stay Rs 1,000 per night, food Rs 500 per day, local Rs 200 per day, "
      f"activities Rs 100 per day, transport Rs 1,000 per person round trip, fees Rs 300 per person, budget style gives a total of "
      f"Rs {E['total']:,.0f} (Rs {E['total']/3:,.0f} per person)."),
    P("All figures are estimates derived from sample data, not live prices."))
sec("20. Security Design",
    B("bcrypt hashing; JWT with expiry; per-user authorisation (404 for foreign ids).", "In-memory AI rate limit; CORS allow-list; secrets only in backend env.",
      "Input validation with Pydantic; generic 500 messages.", "Limitation: JWT in localStorage; no CSRF/cookie mode; no password reset."))
sec("21. Implementation Details",
    P("Backend: routers call services and pure domain functions. Frontend: typed API client, contexts for auth, toasts and motion, `useAsync` for loading/error states. "
      "Diagrams are generated by docs/make_diagrams.py; documents by the build scripts in docs/."))
sec("22. Testing Methodology",
    P("Backend tests use pytest and FastAPI's TestClient with a temporary SQLite database, LLM disabled and live weather/currency disabled. "
      "Budget and travel logic are tested as pure functions with hand-calculated values. Frontend helpers are tested with Vitest. A manual checklist is in docs/TEST_PLAN.md."))
sec("23. Test Cases", P(f"{len(tests)} automated test functions exist in the repository:"), T(["File", "Test function"], tests))
sec("24. Validation and Error Handling",
    B("Pydantic rejects bad types, ranges and date order (HTTP 422).", "Duplicate email returns 409; wrong credentials return 401; foreign trips return 404.",
      "Unsupported start cities and over-budget plans produce clear messages and alternatives.", "LLM, weather and currency failures degrade to labelled fallbacks."))
sec("25. Actual Test Results",
    P("No test run has been recorded in this report. Run the commands below and paste the real output here."),
    ("code", "cd backend && pytest -v\ncd frontend && npm run typecheck && npm test"),
    PH("[Paste pytest output here]"), PH("[Paste vitest output here]"), PH("[Number passed / failed, date, environment]"))
sec("26. Limitations", B(*c.LIMITATIONS))
sec("27. Future Enhancements", B(*c.FUTURE))
sec("28. Conclusion",
    P("TourPlan AI shows how a deterministic, testable planning core can be combined with an optional language model without letting the model "
      "invent facts. The design favours transparency (visible scores, labelled estimates) over impressive but unverifiable output."))
sec("29. References", B(*c.REFERENCES))
sec("30. Appendix A: Setup Commands",
    ("code", "cd backend\npython -m venv .venv\n# activate the venv\npip install -r requirements.txt\npytest\nuvicorn app.main:app --reload\n\n"
             "cd frontend\nnpm install\nnpm run typecheck\nnpm test\nnpm run dev"))
sec("31. Appendix B: Sample Data Notice",
    P("All destinations, attractions, prices and durations are hand-written sample data for demonstration. Prices are indicative estimates. "
      "Coordinates are approximate. The application does not book travel."))

# ------------------------------------------------------------------ DOCX
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, RGBColor


def field(par, instr):
    for t, txt in (("begin", None), (None, instr), ("end", None)):
        r = par.add_run()
        if t:
            e = OxmlElement("w:fldChar")
            e.set(qn("w:fldCharType"), t)
        else:
            e = OxmlElement("w:instrText")
            e.set(qn("xml:space"), "preserve")
            e.text = txt
        r._r.append(e)


def build_docx(out):
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    for s in d.sections:
        s.left_margin = s.right_margin = Inches(1)
        fp = s.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.add_run(c.PROJECT + "  |  Page ")
        field(fp, "PAGE")
    for i, (title, blocks) in enumerate(S):
        if i < 3 or title == "Abstract":
            pass
        h = d.add_heading(title, level=1)
        if i == 0:
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for b in blocks:
            k = b[0]
            if k == "p":
                p = d.add_paragraph(b[1])
                if i == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif k == "ph":
                p = d.add_paragraph()
                r = p.add_run(b[1])
                r.font.highlight_color = 7  # yellow: editable placeholder
                if i == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif k == "b":
                for x in b[1]:
                    d.add_paragraph(x, style="List Bullet")
            elif k == "code":
                p = d.add_paragraph()
                r = p.add_run(b[1])
                r.font.name = "Consolas"
                r.font.size = Pt(9)
            elif k == "t":
                t = d.add_table(rows=1, cols=len(b[1]))
                t.style = "Light Grid Accent 1"
                for j, hd in enumerate(b[1]):
                    t.rows[0].cells[j].text = hd
                    for r in t.rows[0].cells[j].paragraphs[0].runs:
                        r.font.bold = True
                for row in b[2]:
                    cells = t.add_row().cells
                    for j, v in enumerate(row):
                        cells[j].text = str(v)
                        for pp in cells[j].paragraphs:
                            for r in pp.runs:
                                r.font.size = Pt(9)
                d.add_paragraph()
            elif k == "img":
                d.add_picture(str(b[1]), width=Inches(6.2))
                cp = d.add_paragraph(b[2])
                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cp.runs[0].italic = True
            elif k == "toc":
                p = d.add_paragraph()
                field(p, 'TOC \\o "1-1" \\h \\z \\u')
                d.add_paragraph("(In Word: right-click the table of contents and choose Update Field to fill it.)").runs[0].italic = True
        if i in (0, 1, 2, 3, 4):
            d.add_page_break()
    d.save(out)


# ------------------------------------------------------------------ PDF
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (BaseDocTemplate, Frame, Image, PageBreak, PageTemplate, Paragraph, Preformatted,
                                Spacer, Table, TableStyle)
from reportlab.platypus.tableofcontents import TableOfContents
from xml.sax.saxutils import escape


class Doc(BaseDocTemplate):
    def afterFlowable(self, f):
        if isinstance(f, Paragraph) and f.style.name == "H1":
            t = f.getPlainText()
            key = "h_" + "".join(ch for ch in t if ch.isalnum())
            self.canv.bookmarkPage(key)
            self.notify("TOCEntry", (0, t, self.page, key))


def build_pdf(out):
    ss = getSampleStyleSheet()
    H1 = ParagraphStyle("H1", parent=ss["Heading1"], fontSize=17, spaceAfter=10, textColor=colors.HexColor("#14532d"))
    N = ParagraphStyle("N", parent=ss["BodyText"], fontSize=10.5, leading=15, spaceAfter=6)
    CEN = ParagraphStyle("CEN", parent=N, alignment=1, fontSize=13)
    SM = ParagraphStyle("SM", parent=N, fontSize=8, leading=10, spaceAfter=0)
    CAP = ParagraphStyle("CAP", parent=N, alignment=1, fontSize=9, textColor=colors.grey)
    PHS = ParagraphStyle("PH", parent=N, backColor=colors.HexColor("#fff3b0"), borderPadding=3)
    doc = Doc(str(out), pagesize=A4, leftMargin=2.2 * cm, rightMargin=2.2 * cm, topMargin=2 * cm, bottomMargin=2 * cm,
              title=c.PROJECT)

    def footer(cv, d):
        cv.setFont("Helvetica", 8)
        cv.drawCentredString(A4[0] / 2, 1 * cm, f"{c.PROJECT}  |  Page {d.page}")

    doc.addPageTemplates([PageTemplate(id="p", frames=[Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height)], onPage=footer)])
    story = []
    for i, (title, blocks) in enumerate(S):
        story.append(Paragraph(escape(title), H1))
        for b in blocks:
            k = b[0]
            if k == "p":
                story.append(Paragraph(escape(b[1]), CEN if i == 0 else N))
            elif k == "ph":
                story.append(Paragraph(escape(b[1]), PHS))
                story.append(Spacer(1, 4))
            elif k == "b":
                for x in b[1]:
                    story.append(Paragraph("&bull; " + escape(x), N))
            elif k == "code":
                story.append(Preformatted(b[1], ParagraphStyle("c", fontName="Courier", fontSize=8.5, leading=11,
                                                               backColor=colors.HexColor("#f0f0f0"))))
                story.append(Spacer(1, 6))
            elif k == "t":
                rows = [[Paragraph("<b>%s</b>" % escape(h), SM) for h in b[1]]] + \
                       [[Paragraph(escape(str(v)), SM) for v in r] for r in b[2]]
                n = len(b[1])
                widths = {2: [3.2 * cm, doc.width - 3.2 * cm], 3: [doc.width * .25, doc.width * .4, doc.width * .35]}.get(n)
                if n == 2 and b[1][0] == "File":
                    widths = [4.5 * cm, doc.width - 4.5 * cm]
                tb = Table(rows, colWidths=widths, repeatRows=1)
                tb.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), .4, colors.grey), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#d9ead3")),
                                        ("VALIGN", (0, 0), (-1, -1), "TOP")]))
                story += [tb, Spacer(1, 8)]
            elif k == "img":
                from PIL import Image as PI
                w, h = PI.open(b[1]).size
                W = doc.width
                story += [Image(str(b[1]), width=W, height=W * h / w), Paragraph(escape(b[2]), CAP)]
            elif k == "toc":
                toc = TableOfContents()
                toc.levelStyles = [ParagraphStyle("t0", fontSize=10.5, leading=15)]
                story.append(toc)
        story.append(PageBreak())
    doc.multiBuild(story)


if __name__ == "__main__":
    (ROOT / "report").mkdir(exist_ok=True)
    build_docx(ROOT / "report/TourPlan_AI_Report.docx")
    build_pdf(ROOT / "report/TourPlan_AI_Report.pdf")
    print("sections:", len(S) - 6 + 0, "(numbered 1-31 + front matter)")
